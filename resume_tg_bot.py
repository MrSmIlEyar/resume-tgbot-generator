import os
from aiogram import Bot, types
from aiogram.dispatcher import Dispatcher
from aiogram.types import ParseMode
from aiogram.utils import executor
from docx import Document

API_TOKEN = '6137850721:AAFSamidQjOex0cjEF5G1SYqtj1knbeZHz8'
placeholders = ['{NAME}', '{AGE}', '{EMAIL}', '{PHONE}', '{EXPERIENCE}', '{EDUCATION}', '{SKILLS}']

bot = Bot(token=API_TOKEN)
dp = Dispatcher(bot)

@dp.message_handler(commands=['start', 'help'])
async def send_welcome(message: types.Message):
    await message.reply("Привет! Я бот для создания резюме. Введите свои данные в следующем формате:\n\n"
                        "ФИО: Ваши Фамилия Имя Отчество\n"
                        "Возраст: Ваш возраст\n"
                        "Email: Ваш email\n"
                        "Телефон: Ваш номер телефона\n"
                        "Опыт работы: Ваш опыт работы\n"
                        "Образование: Ваше образование\n"
                        "Навыки: Ваши навыки")

@dp.message_handler()
async def fill_resume_template(message: types.Message):
    user_data = message.text.split('\n')

    if len(user_data) != len(placeholders):
        await message.reply('Пожалуйста, предоставьте информацию в правильном формате. Используйте команду /help для помощи.')
        return

    doc = Document('Documents/blank-rezume-1.docx')

    for paragraph in doc.paragraphs:
        for i, placeholder in enumerate(placeholders):
            if placeholder in paragraph.text:
                paragraph.text = paragraph.text.replace(placeholder, user_data[i])

    doc.save(f'resume_{message.from_user.id}.docx')
    with open(f'resume_{message.from_user.id}.docx', 'rb') as resume_file:
        await bot.send_document(chat_id=message.chat.id, document=resume_file)
        resume_file.close()

    os.remove(f'resume_{message.from_user.id}.docx')

if __name__ == '__main__':
    from aiogram import executor
    executor.start_polling(dp, skip_updates=True)